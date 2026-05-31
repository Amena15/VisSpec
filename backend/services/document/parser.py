"""Document parsing service for PDF, DOCX, and TXT files."""
import io
import re
from typing import Dict, Any, Optional, List, Tuple

try:
    from services.ai.text_cleaner import TextCleaner
except ImportError:
    TextCleaner = None

# Import enhanced parser
from services.document.parser_enhanced import EnhancedDocumentParser


def _collapse_repeated_word_token(match: re.Match) -> str:
    """Collapse 'willwill' -> 'will' but never touch digit runs like '000' in '25,000'."""
    word = match.group(1)
    if word.isdigit():
        return match.group(0)
    return word


class DocumentParser:
    """Parses PDF and DOCX documents to extract text, structure, and requirements."""

    # Common SRS section headers for structure identification (REQ_06)
    SECTION_PATTERNS = [
        r"^(?:\d+\.)+\s*(?:Functional Requirements|System Overview|Introduction|Scope|Purpose)",
        r"^(?:\d+\.)+\s*(?:Non-Functional Requirements|Performance Requirements|Security Requirements)",
        r"^(?:\d+\.)+\s*(?:Interface Requirements|User Interface|External Interface)",
        r"^(?:\d+\.)+\s*(?:System Features|Use Cases|User Requirements)",
        r"^(?:\d+\.)+\s*(?:Constraints|Assumptions|Dependencies)",
        r"^(?:\d+\.)+\s*(?:Data Requirements|Database|Data Model)",
        r"^(?:\d+\.)+\s*(?:Overall Description|Product Perspective|Product Features)",
        r"^(?:\d+\.)+\s*(?:Specific Requirements|Functional Requirements|User Requirements)",
        r"^(?:\d+\.)+\s*(?:Other Non-Functional Requirements|Performance Requirements)",
        r"^(?:\d+\.)+\s*(?:Availability Requirements|Reliability|Security|Usability)",
        r"^(?:\d+\.)+\s*(?:Scalability|Maintainability|Portability|Backup)",
        r"^(?:\d+\.)+\s*(?:Business Rules|Summary|Appendices)",
        r"^\d+\.\d+\s+",  # Generic numbered section like 1.1, 3.2
        r"^(?:Chapter|Section)\s+\d+",
        r"^(?:APPENDIX|Appendix)\s+[A-Z]",
        r"^\d+\.\s*[A-Z][a-zA-Z\s]+",  # Main sections like "1. Introduction"
        r"^\d+\.\d+\s*[A-Z][a-zA-Z\s]+",  # Subsections like "2.1 Product Perspective"
    ]
    
    # Fallback patterns for domain-specific or custom section headers
    SECTION_FALLBACK_PATTERNS = [
        r'^[A-Z][A-Z\s]+:$',  # All caps ending with colon (e.g., "INTRODUCTION:")
        r'^[A-Z][a-zA-Z\s]+:$',  # Title case ending with colon (e.g., "Functional Requirements:")
        r'^[A-Z]{3,}$',  # All caps (e.g., "INTRODUCTION")
        r'^\d+\.\s+[A-Z][a-zA-Z]+$',  # Numbered with single word (e.g., "1. Introduction")
    ]

    # Requirement ID patterns (REQ_05)
    REQ_ID_PATTERNS = [
        r"(REQ[_\-]?\d+)",
        r"(FR[_\-]?\d+)",
        r"(NFR[_\-]?\d+)",
        r"(SR[_\-]?\d+)",
        r"(\d+\.\d+\.\d+)",
        r"(UC[_\-]?\d+)",
    ]

    def parse_pdf(self, file_content: bytes, filename: str) -> Dict[str, Any]:
        """Parse a PDF file and extract text with structure and table extraction."""
        import fitz
        import io
        
        doc = fitz.open(stream=io.BytesIO(file_content), filetype="pdf")
        
        result = {
            "filename": filename,
            "file_extension": "pdf",
            "text": "",
            "page_count": len(doc),
            "word_count": 0,
            "is_scanned_pdf": 0,
            "extracted_author": "",
            "extracted_version": "",
            "tables_found": 0,
            "table_rows_extracted": 0,
        }
        
        try:
            # Extract text from each page with structure preservation
            full_text = ""
            
            for page_num in range(len(doc)):
                page = doc[page_num]
                
                # Get text with layout information
                text = page.get_text()
                
                # Add page separator for structure preservation
                if text.strip():
                    full_text += text + "\n\n"
            
            # Extract table content using camelot-py (Phase 2.1)
            table_text = self._extract_pdf_tables(doc)
            if table_text:
                full_text += "\n\n" + table_text
                result["tables_found"] = len([page for page in doc if self._page_has_tables(page)])
                result["table_rows_extracted"] = len(table_text.split('\n'))
                logger.info(f"Extracted content from {result['tables_found']} table pages")
            
            # Check if PDF is scanned (very little text)
            if self._detect_scanned(doc):
                result["is_scanned_pdf"] = 1
                # Try enhanced OCR as fallback (Phase 2.2)
                ocr_text = self._attempt_enhanced_ocr(doc)
                if ocr_text:
                    full_text = ocr_text
                else:
                    result["error"] = "No text detected in scanned PDF. OCR failed. Please use a text-based PDF or convert the document to text format."
            
            # Apply cleanup for merged modal verbs
            cleaned_text = self._cleanup_merged_modals(full_text)
            
            result["extracted_text"] = cleaned_text
            result["text"] = cleaned_text
            result["word_count"] = len(cleaned_text.split())
            
            # Extract metadata
            metadata = doc.metadata
            result["extracted_author"] = metadata.get("author", "")
            result["extracted_version"] = str(metadata.get("creationDate", ""))
            
            doc.close()
            return result
        except ImportError:
            return {"error": "PyMuPDF not installed. Run: pip install PyMuPDF"}
        except Exception as e:
            return {"error": str(e)}

    def parse_docx(self, file_content: bytes, filename: str) -> Dict[str, Any]:
        """Parse a DOCX file and extract text and metadata including table content."""
        try:
            from docx import Document
            doc = Document(io.BytesIO(file_content))
            
            # Extract paragraph text
            paragraph_text = "\n".join(p.text for p in doc.paragraphs)
            
            # Extract table content (this is where functional requirements often are)
            table_text = []
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        cell_text = cell.text.strip()
                        if cell_text:
                            row_text.append(cell_text)
                    if row_text:
                        # Process each cell as a potential requirement
                        for cell_text in row_text:
                            if len(cell_text.strip()) > 5:  # Skip very short cells
                                table_text.append(cell_text.strip())
                        # Also keep the combined row for context
                        table_text.append(" | ".join(row_text))
            
            # Combine paragraph and table text
            all_text = paragraph_text + "\n\n" + "\n".join(table_text)
            
            # Apply basic cleanup for merged modal verbs
            cleaned_text = self._cleanup_merged_modals(all_text)

            return {
                "filename": filename,
                "file_extension": "docx",
                "text": cleaned_text,
                "page_count": None,  # DOCX doesn't have fixed pages
                "word_count": len(cleaned_text.split()),
                "is_scanned_pdf": 0,
                "extracted_author": doc.core_properties.author or "",
                "extracted_version": str(doc.core_properties.revision or ""),
                "tables_found": len(doc.tables),
                "table_rows_extracted": sum(len(table.rows) for table in doc.tables),
            }
        except ImportError:
            return {"error": "python-docx not installed. Run: pip install python-docx"}
        except Exception as e:
            return {"error": str(e)}

    def parse_txt(self, file_content, filename: str) -> Dict[str, Any]:
        """Parse a plain text file."""
        # Handle both bytes and string input
        if isinstance(file_content, bytes):
            text = file_content.decode("utf-8", errors="replace")
        else:
            text = str(file_content)
        
        return {
            "filename": filename,
            "file_extension": "txt",
            "text": text,
            "page_count": 1,
            "word_count": len(text.split()),
            "is_scanned_pdf": 0,
            "extracted_author": "",
            "extracted_version": "",
        }

    def parse(self, file_content: bytes, filename: str) -> Dict[str, Any]:
        """Auto-detect format and parse the document."""
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        if ext == "pdf":
            return self.parse_pdf(file_content, filename)
        elif ext == "docx":
            return self.parse_docx(file_content, filename)
        elif ext == "txt":
            return self.parse_txt(file_content, filename)
        else:
            return {"error": f"Unsupported file format: .{ext}"}

    @staticmethod
    def _detect_scanned(pdf_doc) -> int:
        """Detect if a PDF is scanned (contains images but little text)."""
        total_text = sum(len(page.get_text().strip()) for page in pdf_doc)
        return 1 if total_text < 100 else 0
    
    def _cleanup_merged_modals(self, text: str) -> str:
        """Clean up merged modal verbs like 'shallshould', 'shallwill', etc."""
        if not text:
            return text
        
        # Common merged modal patterns
        merged_patterns = {
            'shallshould': 'shall',
            'shallwill': 'shall',
            'shallmust': 'shall',
            'shouldwill': 'should',
            'shouldmust': 'should',
            'willmust': 'will',
            'mayshall': 'may',
            'canshall': 'can',
            'mustshould': 'must',
            'mustwill': 'must',
        }
        
        # Apply replacements
        cleaned_text = text
        for merged, correct in merged_patterns.items():
            cleaned_text = cleaned_text.replace(merged, correct)
        
        # Only remove exact duplicates with no space in between (like 'shallshall')
        # Don't remove legitimate modal verbs
        cleaned_text = re.sub(r'(shall|must|should|will|may|can)\1', r'\1', cleaned_text)
        
        return cleaned_text
    
    def _extract_pdf_tables(self, pdf_doc) -> str:
        """Extract table content from PDF using camelot-py (Phase 2.1)."""
        try:
            import camelot
            import io
            import pandas as pd
            
            table_text = []
            tables_found = 0
            
            for page_num in range(len(pdf_doc)):
                try:
                    # Extract tables from current page
                    tables = camelot.read_pdf(io.BytesIO(pdf_doc.write()), pages=str(page_num + 1))
                    
                    if tables:
                        tables_found += len(tables)
                        logger.info(f"Found {len(tables)} tables on page {page_num + 1}")
                        
                        for table in tables:
                            # Convert table to structured text
                            table_df = table.df
                            
                            # Process each row
                            for _, row in table_df.iterrows():
                                # Skip empty rows
                                if row.isna().all():
                                    continue
                                
                                # Clean and join cell content
                                row_text = []
                                for cell in row:
                                    if pd.notna(cell) and str(cell).strip():
                                        row_text.append(str(cell).strip())
                                
                                if row_text:
                                    # Join cells with | separator for requirement extraction
                                    table_text.append(" | ".join(row_text))
                    
                except Exception as e:
                    logger.debug(f"Could not extract tables from page {page_num + 1}: {e}")
                    continue
            
            if table_text:
                logger.info(f"Extracted {len(table_text)} table rows from {tables_found} tables")
                return "\n".join(table_text)
            
            return ""
            
        except ImportError:
            logger.warning("camelot-py not installed. PDF table extraction disabled.")
            return ""
        except Exception as e:
            logger.error(f"PDF table extraction failed: {e}")
            return ""

    def _page_has_tables(self, page) -> bool:
        """Check if a PDF page contains table-like structures."""
        try:
            # Simple heuristic: look for repeated tabular patterns
            text = page.get_text()
            lines = text.split('\n')
            
            # Look for lines with multiple tab-separated or space-separated columns
            tabular_lines = 0
            for line in lines:
                if line.strip():
                    # Check for patterns suggesting table structure
                    if (line.count('\t') >= 2 or 
                        len(line.split()) >= 4 or
                        any(char in line for char in ['|', '├', '┤', '┬', '┴'])):
                        tabular_lines += 1
            
            return tabular_lines >= 3  # At least 3 tabular lines suggest a table
            
        except Exception:
            return False

    def _attempt_enhanced_ocr(self, pdf_doc) -> Optional[str]:
        """Enhanced OCR for scanned PDFs with multiple strategies (Phase 2.2)."""
        try:
            import pytesseract
            from PIL import Image
            import io
            
            all_text = []
            ocr_used = False
            
            for page_num, page in enumerate(pdf_doc):
                page_text = ""
                
                # Strategy 1: Try with PSM 6 (uniform block of text)
                try:
                    pix = page.get_pixmap()
                    img_data = pix.tobytes("png")
                    image = Image.open(io.BytesIO(img_data))
                    
                    text = pytesseract.image_to_string(
                        image, 
                        lang='eng',
                        config='--psm 6'
                    )
                    
                    if text.strip() and len(text.strip()) > 20:
                        page_text = text
                        ocr_used = True
                        logger.debug(f"Page {page_num + 1}: OCR PSM 6 successful")
                
                except Exception as e:
                    logger.debug(f"Page {page_num + 1}: OCR PSM 6 failed: {e}")
                
                # Strategy 2: Try with PSM 3 if first attempt failed
                if not page_text.strip():
                    try:
                        pix = page.get_pixmap()
                        img_data = pix.tobytes("png")
                        image = Image.open(io.BytesIO(img_data))
                        
                        text = pytesseract.image_to_string(
                            image,
                            lang='eng', 
                            config='--psm 3'
                        )
                        
                        if text.strip() and len(text.strip()) > 20:
                            page_text = text
                            ocr_used = True
                            logger.debug(f"Page {page_num + 1}: OCR PSM 3 successful")
                    
                    except Exception as e:
                        logger.debug(f"Page {page_num + 1}: OCR PSM 3 failed: {e}")
                
                if page_text.strip():
                    all_text.append(page_text)
            
            if ocr_used:
                combined_text = "\n\n".join(all_text)
                logger.info(f"Enhanced OCR extracted {len(combined_text.split())} words from {len(pdf_doc)} pages")
                return combined_text
            
            return None
            
        except ImportError:
            logger.warning("pytesseract not installed. OCR functionality disabled.")
            return None
        except Exception as e:
            logger.error(f"Enhanced OCR failed: {e}")
            return None

    def _attempt_ocr(self, pdf_doc) -> Optional[str]:
        """Legacy OCR method - delegates to enhanced OCR."""
        return self._attempt_enhanced_ocr(pdf_doc)

    def identify_structure(self, text: str) -> List[Dict[str, Any]]:
        """
        Identify logical sections/headers in the document.
        Implements REQ_06: Structure Identification.
        Uses fallback patterns for domain-specific headers.
        """
        sections = []
        lines = text.split("\n")
        current_section = None
        section_text = []
        seq = 0

        for line in lines:
            line_stripped = line.strip()
            if not line_stripped:
                continue

            is_header = False
            
            # Try primary section patterns first
            for pattern in self.SECTION_PATTERNS:
                if re.match(pattern, line_stripped, re.IGNORECASE):
                    # Save previous section
                    if current_section:
                        sections.append({
                            "title": current_section,
                            "text": "\n".join(section_text).strip(),
                            "sequence_order": seq,
                        })
                    current_section = line_stripped
                    section_text = []
                    seq += 1
                    is_header = True
                    break
            
            # If not matched by primary patterns, try fallback patterns
            if not is_header:
                for pattern in self.SECTION_FALLBACK_PATTERNS:
                    if re.match(pattern, line_stripped, re.IGNORECASE):
                        # Save previous section
                        if current_section:
                            sections.append({
                                "title": current_section,
                                "text": "\n".join(section_text).strip(),
                                "sequence_order": seq,
                            })
                        current_section = line_stripped
                        section_text = []
                        seq += 1
                        is_header = True
                        break

            if not is_header and current_section:
                section_text.append(line_stripped)
            elif not is_header and not current_section:
                # Text before first header
                current_section = "Introduction"
                section_text.append(line_stripped)

        # Save last section
        if current_section:
            sections.append({
                "title": current_section,
                "text": "\n".join(section_text).strip(),
                "sequence_order": seq,
            })

        return sections

    def extract_requirements(self, text: str) -> List[Dict[str, Any]]:
        """
        Enhanced requirement extraction for professional quality.
        Focuses on complete sentences and proper requirement identification.
        """
        if not text:
            return []
        
        # Use enhanced parser for better quality
        enhanced_parser = EnhancedDocumentParser()
        
        # Clean text first
        cleaned_text = self.sanitize_text(text)
        
        # Extract requirements with enhanced quality
        requirements = enhanced_parser.extract_requirements(cleaned_text)
        
        return requirements
    
    def _finalize_requirement(self, req_data: Dict[str, Any], requirements: List[Dict[str, Any]], seq: int):
        """Finalize and add a requirement to the list."""
        cleaned_text = re.sub(r'\s+', ' ', req_data["text"])  # Normalize whitespace
        cleaned_text = re.sub(r'\.+$', '', cleaned_text)  # Remove trailing period
        
        # Try to find requirement ID
        req_id = None
        for pattern in self.REQ_ID_PATTERNS:
            match = re.search(pattern, cleaned_text, re.IGNORECASE)
            if match:
                req_id = match.group(1).upper()
                break

        requirements.append({
            "req_identifier": req_id or f"REQ-{seq+1:03d}",
            "original_text": cleaned_text,
            "word_count": len(cleaned_text.split()),
            "sequence_order": seq,
        })

    def sanitize_text(self, text: str) -> str:
        """Sanitize extracted text for storage and search."""
        if not text:
            return ""
        
        # Remove duplicate consecutive words with space (conservative approach)
        # Only matches "will will" not "willwill" to avoid false positives
        # Also catches triple repeats (will will will -> will)
        text = re.sub(r'\b(\w+)(?:\s+\1)+\b', r'\1', text, flags=re.IGNORECASE)
        
        # Modal verbs jammed without space (e.g. shallwill) — \b between modals does not match
        try:
            text = TextCleaner.fix_concatenated_modals(text)
        except Exception as e:
            # If TextCleaner fails, use simple regex fallback
            text = re.sub(r'(shall|must|should|will|may|can)(shall|must|should|will|may|can)', r'\1', text, flags=re.IGNORECASE)
        
        # Handle concatenated duplicates (willwill -> will); skip digit-only tokens (preserves 25,000)
        text = re.sub(r'\b(\w+)(\1)+\b', _collapse_repeated_word_token, text, flags=re.IGNORECASE)
        
        # Remove page numbers (but preserve section headers)
        # Page numbers are typically standalone or at the end of lines
        text = re.sub(r'^\s*\d+\s*$', '', text, flags=re.MULTILINE)  # Standalone numbers
        text = re.sub(r'Page\s+\d+\s+of\s+\d+', '', text, flags=re.IGNORECASE)  # "Page 1 of 3"
        text = re.sub(r'\n\s*\d+\s*\n', '\n', text)  # Numbers alone on lines
        
        # Remove multiple newlines but preserve single newlines
        text = re.sub(r'\n\s*\n', '\n', text)
        
        # Remove extra spaces but preserve newlines
        text = re.sub(r'[ \t]+', ' ', text)  # Only spaces and tabs, not newlines
        
        return text.strip()

    def parse_full(self, file_content: bytes, filename: str) -> Dict[str, Any]:
        """
        Full parsing pipeline: extract text, sanitize, identify structure,
        extract requirements.
        """
        # Step 1: Basic parsing
        result = self.parse(file_content, filename)
        if "error" in result:
            return result

        text = result.get("text", "")

        # Step 2: Sanitize (remove noise)
        cleaned_text = self.sanitize_text(text)
        result["extracted_text"] = cleaned_text

        # Step 3: Structure identification (REQ_06)
        sections = self.identify_structure(cleaned_text)
        result["sections"] = sections
        result["sections_found"] = [s["title"] for s in sections]

        # Step 4: Requirement extraction (REQ_05)
        requirements = self.extract_requirements(cleaned_text)
        result["requirements"] = requirements
        result["requirements_found"] = len(requirements)

        return result

